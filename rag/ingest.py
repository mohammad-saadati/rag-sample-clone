"""Stage 1 - Ingestion: turn PDFs, images, Office/HTML/text files into Documents.

Each Document keeps page-level text (so answers can cite pages) and metadata
(source, type, year, plus anything from an optional `<file>.meta.json` sidecar).
"""
from __future__ import annotations

import hashlib
import json
import logging
import os
import re
from concurrent.futures import ProcessPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from .text import normalize

log = logging.getLogger(__name__)

TEXT_EXT = {".txt", ".md", ".rst", ".csv", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED = TEXT_EXT | IMAGE_EXT | {".pdf", ".docx", ".html", ".htm"}


@dataclass
class Document:
    doc_id: str
    pages: list[str]                      # page texts (a text file = 1 page)
    metadata: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)


# ---------- per-format readers ----------

def _read_pdf(path: Path, ocr_lang: str) -> list[str]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        if len(txt.strip()) < 20:          # likely a scanned page -> try OCR
            txt = _ocr_pdf_page(path, i, ocr_lang) or txt
        pages.append(txt)
    return pages


def _ocr_pdf_page(path: Path, index: int, lang: str) -> str:
    try:
        from pdf2image import convert_from_path
        import pytesseract
        imgs = convert_from_path(str(path), first_page=index + 1, last_page=index + 1, dpi=200)
        return pytesseract.image_to_string(imgs[0], lang=lang) if imgs else ""
    except Exception as e:  # optional dependency or poppler missing
        log.debug("OCR for %s p%d skipped: %s", path, index + 1, e)
        return ""


def _read_image(path: Path, lang: str) -> list[str]:
    try:
        import pytesseract
        from PIL import Image
        return [pytesseract.image_to_string(Image.open(path), lang=lang)]
    except Exception as e:
        log.warning("Skipping image %s (install pytesseract + tesseract for OCR): %s", path, e)
        return []


def _read_docx(path: Path) -> list[str]:
    import docx  # python-docx
    d = docx.Document(str(path))
    return ["\n".join(p.text for p in d.paragraphs)]


def _read_html(path: Path) -> list[str]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    import html
    return [html.unescape(raw)]


def _read_text(path: Path) -> list[str]:
    return [path.read_text(encoding="utf-8", errors="ignore")]


# ---------- public API ----------

def load_file(path: str | Path, ocr_lang: str = "eng+fas") -> Document | None:
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            pages = _read_pdf(path, ocr_lang)
        elif ext in IMAGE_EXT:
            pages = _read_image(path, ocr_lang)
        elif ext == ".docx":
            pages = _read_docx(path)
        elif ext in {".html", ".htm"}:
            pages = _read_html(path)
        elif ext in TEXT_EXT:
            pages = _read_text(path)
        else:
            return None
    except Exception as e:
        log.warning("Failed to read %s: %s", path, e)
        return None

    pages = [normalize(p) for p in pages]
    if not any(p.strip() for p in pages):
        return None

    stat = path.stat()
    modified = datetime.fromtimestamp(stat.st_mtime)
    meta = {
        "source": str(path),
        "filename": path.name,
        "title": path.stem.replace("_", " ").replace("-", " "),
        "type": ext.lstrip("."),
        "num_pages": len(pages),
        "year": modified.year,
        "modified": modified.isoformat(timespec="seconds"),
    }
    sidecar = path.with_name(path.name + ".meta.json")
    if sidecar.exists():
        try:
            meta.update(json.loads(sidecar.read_text(encoding="utf-8")))
        except json.JSONDecodeError as e:
            log.warning("Bad sidecar %s: %s", sidecar, e)

    content_hash = hashlib.sha1("".join(pages).encode("utf-8")).hexdigest()
    meta["content_hash"] = content_hash
    doc_id = hashlib.sha1(str(path.resolve()).encode("utf-8")).hexdigest()[:16]
    return Document(doc_id=doc_id, pages=pages, metadata=meta)


def iter_paths(paths: list[str | Path]):
    for p in paths:
        p = Path(p)
        if p.is_dir():
            for root, _, files in os.walk(p):
                for f in sorted(files):
                    fp = Path(root) / f
                    if fp.suffix.lower() in SUPPORTED and not f.endswith(".meta.json"):
                        yield fp
        elif p.suffix.lower() in SUPPORTED:
            yield p


def load_documents(paths, workers: int = 1, ocr_lang: str = "eng+fas") -> list[Document]:
    """Load every supported file. Use workers>1 for large corpora (100k+ files)."""
    files = list(iter_paths(paths))
    if workers > 1 and len(files) > 50:
        with ProcessPoolExecutor(max_workers=workers) as ex:
            docs = list(ex.map(load_file, files, [ocr_lang] * len(files), chunksize=32))
    else:
        docs = [load_file(f, ocr_lang) for f in files]

    seen, out = set(), []
    for d in docs:                        # exact-duplicate removal by content hash
        if d and d.metadata["content_hash"] not in seen:
            seen.add(d.metadata["content_hash"])
            out.append(d)
    log.info("Loaded %d documents (%d files scanned)", len(out), len(files))
    return out
